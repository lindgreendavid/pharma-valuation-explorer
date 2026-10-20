"""Build the downloadable Excel model with live formulas.

Sheets: README, Data (every fact with its source), Comparison (formulas over Data), DCF (inputs, projection,
terminal value, enterprise-to-equity bridge), Sensitivity (formulas, one per cell) and Checks.
Cached values are not stored: Excel and LibreOffice calculate on open, and tests/test_workbook.py
recalculates the file independently and compares it with the Python model.
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parents[1]
sys.path.insert(0, str(ROOT / "src"))

from pharma_valuation.defaults import (  # noqa: E402
    BASE_YEAR,
    BETA,
    DEBT_WEIGHT,
    GROWTH,
    OPERATING_MARGIN,
    TAX_RATE,
    TERMINAL_GROWTH,
    market_parameters,
)
from pharma_valuation.model import Dataset  # noqa: E402

OUT = ROOT / "site" / "downloads" / "pharma-valuation-model.xlsx"
INPUT_FONT = Font(color="0000FF")
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="E1EEE6")
COMPANIES = ["Novartis", "GSK", "Novo Nordisk"]
YEARS = [2023, 2024, 2025]


def sumifs(rows: int, company: str, year_ref: str | int, field: str) -> str:
    year = year_ref if isinstance(year_ref, str) else str(year_ref)
    return (
        f'SUMIFS(Data!$D$2:$D${rows + 1},Data!$A$2:$A${rows + 1},"{company}",'
        f'Data!$B$2:$B${rows + 1},{year},Data!$C$2:$C${rows + 1},"{field}")'
    )


def build() -> Path:
    data = Dataset.load()
    facts = data.payload["facts"]
    params = market_parameters()
    wb = Workbook()

    readme = wb.active
    readme.title = "README"
    lines = [
        "Pharma Valuation Explorer: comparison and DCF model (version 0.1.0)",
        "",
        "What this is: public IFRS financials of Novartis (USD), GSK (GBP) and Novo Nordisk (DKK) for fiscal years 2023 to 2025, a comparison sheet, and a DCF of Novartis.",
        "Source of every number on the Data sheet: the companies' Form 20-F filings on SEC EDGAR (accession number and R page in each row). Millions of reporting currency, as filed.",
        "Fiscal years end 31 December for all three companies. Figures are as reported under IFRS. Company 'core' or 'adjusted' measures are not used.",
        "Novartis figures are continuing operations (the 2023 Sandoz spin-off is excluded). Revenue is net sales; other revenue and royalty income are excluded from revenue.",
        "Blue cells are inputs you can change. Everything else is a formula.",
        "The DCF defaults are the author's assumptions for exploring a range. They are not forecasts, guidance or investment advice.",
        "No share prices are included: no public price source with a clear reuse licence was available programmatically. The model gives a value per share you can compare with a price you choose.",
        "Risk-free rate: US Treasury 10-year yield on 2025-12-31 (FRED DGS10). Equity risk premium: Damodaran implied ERP, 2026-01-01.",
        "Not covered in the enterprise-to-equity bridge: litigation and tax provisions, off-balance-sheet commitments, pension surpluses, and operating-lease commitments outside IFRS 16 liabilities.",
        "See docs/methodology.md in the repository for definitions, the comparability rules and the validation results.",
    ]
    for index, line in enumerate(lines, start=1):
        readme.cell(row=index, column=1, value=line)
    readme["A1"].font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 160

    ws = wb.create_sheet("Data")
    headers = ["company", "fiscal_year", "field", "value", "currency", "accession", "source_page", "line_label", "note"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = BOLD
        cell.fill = HEAD_FILL
    for fact in facts:
        ws.append(
            [
                fact["company"],
                fact["fiscal_year"],
                fact["field"],
                fact["value"],
                fact["currency"],
                fact["accession"],
                fact["source_page"],
                fact["line_label"],
                fact["note"],
            ]
        )
    n = len(facts)
    for column, width in zip("ABCDEFGHI", (14, 11, 34, 14, 9, 22, 14, 60, 60), strict=True):
        ws.column_dimensions[column].width = width

    # ---------------------------------------------------------- Comparison
    cmp_ws = wb.create_sheet("Comparison")
    cmp_ws["A1"] = "Comparable metrics, as reported (millions of reporting currency)"
    cmp_ws["A1"].font = Font(bold=True, size=13)
    cmp_ws.append([])
    header_row = 3
    cmp_ws.cell(row=header_row, column=1, value="Metric")
    columns: dict[tuple[str, int], str] = {}
    col = 2
    for company in COMPANIES:
        for year in YEARS:
            columns[(company, year)] = get_column_letter(col)
            cmp_ws.cell(row=header_row, column=col, value=f"{company} {year}")
            cmp_ws.cell(row=header_row - 1, column=col, value=data.currency(company))
            col += 1
    for cell in cmp_ws[header_row]:
        cell.font = BOLD
        cell.fill = HEAD_FILL

    row_of: dict[str, int] = {}
    base_fields = [
        ("revenue", "Revenue (net sales)"),
        ("operating_income", "Operating income (reported)"),
        ("r_and_d", "Research and development expense"),
        ("income_tax", "Income tax expense"),
        ("pretax_income", "Income before tax"),
        ("net_income", "Net income (continuing operations)"),
        ("cfo", "Operating cash flow"),
        ("capex_ppe", "Capex: property, plant and equipment"),
        ("capex_intangibles", "Capex: intangible assets"),
        ("interest_paid_outside_cfo", "Interest paid, if reported outside operating cash flow"),
        ("interest_received_outside_cfo", "Interest received, if reported outside operating cash flow"),
        ("debt_current", "Borrowings: current"),
        ("debt_noncurrent", "Borrowings: non-current"),
        ("lease_current", "Lease liabilities: current (if separate)"),
        ("lease_noncurrent", "Lease liabilities: non-current (if separate)"),
        ("contingent_consideration_current", "Contingent consideration: current (debt-like)"),
        ("contingent_consideration_noncurrent", "Contingent consideration: non-current (debt-like)"),
        ("cash", "Cash and cash equivalents"),
        ("short_term_investments", "Marketable securities and short-term investments"),
    ]
    current = header_row + 1
    for field, label in base_fields:
        row_of[field] = current
        cmp_ws.cell(row=current, column=1, value=label)
        for (company, year), letter in columns.items():
            cmp_ws[f"{letter}{current}"] = "=" + sumifs(n, company, year, field)
        current += 1
    current += 1
    derived = [
        (
            "growth",
            "Revenue growth",
            lambda c, y, L: (
                f'=IF(ISNUMBER({columns.get((c, y - 1), "ZZ")}{row_of["revenue"]}),{L}{row_of["revenue"]}/{columns[(c, y - 1)]}{row_of["revenue"]}-1,"")'
                if (c, y - 1) in columns
                else '=""'
            ),
        ),
        ("op_margin", "Operating margin", lambda c, y, L: f"={L}{row_of['operating_income']}/{L}{row_of['revenue']}"),
        ("rd_intensity", "R&D as % of revenue", lambda c, y, L: f"={L}{row_of['r_and_d']}/{L}{row_of['revenue']}"),
        ("tax_rate", "Effective tax rate", lambda c, y, L: f"={L}{row_of['income_tax']}/{L}{row_of['pretax_income']}"),
        ("net_margin", "Net margin", lambda c, y, L: f"={L}{row_of['net_income']}/{L}{row_of['revenue']}"),
        (
            "fcf",
            "Free cash flow (after tax and net interest; before leases, acquisitions, dividends)",
            lambda c, y, L: (
                f"={L}{row_of['cfo']}-{L}{row_of['capex_ppe']}-{L}{row_of['capex_intangibles']}-{L}{row_of['interest_paid_outside_cfo']}+{L}{row_of['interest_received_outside_cfo']}"
            ),
        ),
        ("fcf_margin", "Free cash flow margin", None),
        ("fcf_conv", "Free cash flow / net income", None),
        (
            "gross_debt",
            "Gross debt incl. leases",
            lambda c, y, L: (
                f"={L}{row_of['debt_current']}+{L}{row_of['debt_noncurrent']}+{L}{row_of['lease_current']}+{L}{row_of['lease_noncurrent']}"
            ),
        ),
        ("net_debt", "Net debt incl. debt-like contingent consideration", None),
        ("leverage", "Net debt / operating income", None),
    ]
    for key, label, fn in derived:
        row_of[key] = current
        cmp_ws.cell(row=current, column=1, value=label)
        for (company, year), letter in columns.items():
            if key == "fcf_margin":
                formula = f"={letter}{row_of['fcf']}/{letter}{row_of['revenue']}"
            elif key == "fcf_conv":
                formula = f"={letter}{row_of['fcf']}/{letter}{row_of['net_income']}"
            elif key == "net_debt":
                formula = (
                    f"={letter}{row_of['gross_debt']}+{letter}{row_of['contingent_consideration_current']}+{letter}{row_of['contingent_consideration_noncurrent']}"
                    f"-{letter}{row_of['cash']}-{letter}{row_of['short_term_investments']}"
                )
            elif key == "leverage":
                formula = f"={letter}{row_of['net_debt']}/{letter}{row_of['operating_income']}"
            else:
                assert fn is not None
                formula = fn(company, year, letter)
            cmp_ws[f"{letter}{current}"] = formula
        current += 1
    cmp_ws.column_dimensions["A"].width = 78
    for letter in columns.values():
        cmp_ws.column_dimensions[letter].width = 16
    for key in ("growth", "op_margin", "rd_intensity", "tax_rate", "net_margin", "fcf_margin"):
        for letter in columns.values():
            cmp_ws[f"{letter}{row_of[key]}"].number_format = "0.0%"
    for letter in columns.values():
        cmp_ws[f"{letter}{row_of['fcf_conv']}"].number_format = "0.00"
        cmp_ws[f"{letter}{row_of['leverage']}"].number_format = "0.00"

    # ---------------------------------------------------------- DCF
    dcf = wb.create_sheet("DCF")
    dcf["A1"] = "Novartis DCF (USD millions unless stated)"
    dcf["A1"].font = Font(bold=True, size=13)
    c = "Novartis"
    y = BASE_YEAR

    def put(
        cell: str, label_cell: str, label: str, value: object, input_cell: bool = False, fmt: str | None = None
    ) -> None:
        dcf[label_cell] = label
        dcf[cell] = value
        if input_cell:
            dcf[cell].font = INPUT_FONT
        if fmt:
            dcf[cell].number_format = fmt

    put("B3", "A3", "Base revenue (FY2025 net sales)", "=" + sumifs(n, c, y, "revenue"))
    put(
        "B4",
        "A4",
        "Base net working capital (inventories + trade receivables - trade payables)",
        "=" + sumifs(n, c, y, "inventories") + "+" + sumifs(n, c, y, "receivables") + "-" + sumifs(n, c, y, "payables"),
    )
    put("B5", "A5", "Tax rate (assumption)", TAX_RATE, True, "0.0%")
    put(
        "B6",
        "A6",
        "D&A as % of revenue (FY2025 depreciation + amortisation; excludes impairments)",
        f"=({sumifs(n, c, y, 'depreciation_ppe')}+{sumifs(n, c, y, 'amortisation_intangibles')})/B3",
        False,
        "0.00%",
    )
    put(
        "B7",
        "A7",
        "Capex as % of revenue (FY2025 PP&E + intangibles)",
        f"=({sumifs(n, c, y, 'capex_ppe')}+{sumifs(n, c, y, 'capex_intangibles')})/B3",
        False,
        "0.00%",
    )
    put("B8", "A8", "Net working capital as % of revenue (held constant)", "=B4/B3", False, "0.00%")
    put("B9", "A9", "Terminal growth (assumption)", TERMINAL_GROWTH, True, "0.00%")
    put("B10", "A10", "Risk-free rate (US 10-year, 2025-12-31)", params["risk_free_rate"]["value"], True, "0.00%")
    put("B11", "A11", "Equity beta (assumption)", BETA, True, "0.00")
    put(
        "B12",
        "A12",
        "Equity risk premium (Damodaran implied, 2026-01-01)",
        params["equity_risk_premium"]["value"],
        True,
        "0.00%",
    )
    put(
        "B13",
        "A13",
        "Pre-tax cost of debt (FY2025 interest expense / average gross debt)",
        f"={sumifs(n, c, y, 'interest_expense')}/(({sumifs(n, c, y, 'debt_current')}+{sumifs(n, c, y, 'debt_noncurrent')}+{sumifs(n, c, y, 'lease_current')}+{sumifs(n, c, y, 'lease_noncurrent')}"
        f"+{sumifs(n, c, y - 1, 'debt_current')}+{sumifs(n, c, y - 1, 'debt_noncurrent')}+{sumifs(n, c, y - 1, 'lease_current')}+{sumifs(n, c, y - 1, 'lease_noncurrent')})/2)",
        False,
        "0.00%",
    )
    put(
        "B14",
        "A14",
        "Debt weight in capital structure (assumption; no market price available)",
        DEBT_WEIGHT,
        True,
        "0.0%",
    )
    put("B15", "A15", "Cost of equity", "=B10+B11*B12", False, "0.00%")
    put("B16", "A16", "After-tax cost of debt", "=B13*(1-B5)", False, "0.00%")
    put("B17", "A17", "WACC (overwrite with a number to test a different rate)", "=(1-B14)*B15+B14*B16", True, "0.00%")
    put("B18", "A18", "Shares outstanding (millions, cover page)", "=" + sumifs(n, c, y, "shares_outstanding_cover"))
    put("B19", "A19", "Mid-year discounting (1 = yes, 0 = no)", 0, True)

    dcf["A21"] = "Projection"
    dcf["A21"].font = BOLD
    labels = {
        22: "Year",
        23: "Revenue growth (assumption)",
        24: "Operating margin on net sales (assumption)",
        25: "Revenue",
        26: "Operating income (EBIT)",
        27: "NOPAT",
        28: "D&A",
        29: "Capex",
        30: "Net working capital",
        31: "Increase in net working capital",
        32: "Free cash flow to the firm",
        33: "Discount factor",
        34: "Present value",
    }
    for r, text in labels.items():
        dcf[f"A{r}"] = text
    dcf["B22"] = 0
    dcf["B25"] = "=B3"
    dcf["B30"] = "=B4"
    for i in range(5):
        letter = get_column_letter(3 + i)
        prev = get_column_letter(2 + i)
        dcf[f"{letter}22"] = i + 1
        dcf[f"{letter}23"] = GROWTH[i]
        dcf[f"{letter}23"].font = INPUT_FONT
        dcf[f"{letter}23"].number_format = "0.0%"
        dcf[f"{letter}24"] = OPERATING_MARGIN[i]
        dcf[f"{letter}24"].font = INPUT_FONT
        dcf[f"{letter}24"].number_format = "0.0%"
        dcf[f"{letter}25"] = f"={prev}25*(1+{letter}23)"
        dcf[f"{letter}26"] = f"={letter}25*{letter}24"
        dcf[f"{letter}27"] = f"={letter}26*(1-$B$5)"
        dcf[f"{letter}28"] = f"={letter}25*$B$6"
        dcf[f"{letter}29"] = f"={letter}25*$B$7"
        dcf[f"{letter}30"] = f"={letter}25*$B$8"
        dcf[f"{letter}31"] = f"={letter}30-{prev}30"
        dcf[f"{letter}32"] = f"={letter}27+{letter}28-{letter}29-{letter}31"
        dcf[f"{letter}33"] = f"=1/(1+$B$17)^({letter}22-0.5*$B$19)"
        dcf[f"{letter}34"] = f"={letter}32*{letter}33"
    dcf["A36"] = "Terminal value and enterprise value"
    dcf["A36"].font = BOLD
    dcf["A37"] = "First terminal-year free cash flow (D&A set equal to capex)"
    dcf["B37"] = "=G27*(1+B9)-B8*G25*B9"
    dcf["A38"] = "Terminal value at year 5"
    dcf["B38"] = "=B37/(B17-B9)"
    dcf["A39"] = "Present value of terminal value"
    dcf["B39"] = "=B38*G33"
    dcf["A40"] = "Present value of explicit cash flows"
    dcf["B40"] = "=SUM(C34:G34)"
    dcf["A41"] = "Enterprise value"
    dcf["B41"] = "=B40+B39"
    dcf["A42"] = "Terminal value as % of enterprise value"
    dcf["B42"] = "=B39/B41"
    dcf["B42"].number_format = "0.0%"

    dcf["A44"] = "Enterprise value to equity value (FY2025 balance sheet)"
    dcf["A44"].font = BOLD
    bridge = [
        (45, "Enterprise value", "=B41"),
        (
            46,
            "Less: gross borrowings and lease liabilities",
            "=-("
            + "+".join(
                sumifs(n, c, y, f) for f in ("debt_current", "debt_noncurrent", "lease_current", "lease_noncurrent")
            )
            + ")",
        ),
        (47, "Plus: cash and cash equivalents", "=" + sumifs(n, c, y, "cash")),
        (48, "Plus: marketable securities and short-term investments", "=" + sumifs(n, c, y, "short_term_investments")),
        (49, "Plus: investments in associates (book)", "=" + sumifs(n, c, y, "equity_method_investments")),
        (50, "Plus: other non-current financial assets (book)", "=" + sumifs(n, c, y, "other_financial_assets")),
        (51, "Less: recognised post-employment liabilities", "=-" + sumifs(n, c, y, "pension_liability")),
        (52, "Less: non-controlling interests (book)", "=-" + sumifs(n, c, y, "nci_book")),
    ]
    for r, label, formula in bridge:
        dcf[f"A{r}"] = label
        dcf[f"B{r}"] = formula
    dcf["A53"] = "Equity value"
    dcf["B53"] = "=SUM(B45:B52)"
    dcf["B53"].font = BOLD
    dcf["A54"] = "Net bridge adjustments (equity value minus enterprise value)"
    dcf["B54"] = "=B53-B45"
    dcf["A55"] = "Equity value per share (USD)"
    dcf["B55"] = "=B53/B18"
    dcf["B55"].font = BOLD
    dcf["A56"] = "Market price per share (enter to compare)"
    dcf["B56"].font = INPUT_FONT
    dcf["A57"] = "DCF value vs entered price"
    dcf["B57"] = '=IF(ISNUMBER(B56),B55/B56-1,"enter a price")'
    dcf["B57"].number_format = "0.0%"
    dcf["A59"] = (
        "Not in this bridge: litigation and tax provisions, off-balance-sheet commitments, pension surpluses, operating-lease commitments outside IFRS 16."
    )
    dcf.column_dimensions["A"].width = 88
    for letter in "BCDEFG":
        dcf.column_dimensions[letter].width = 16

    # ---------------------------------------------------------- Sensitivity
    sens = wb.create_sheet("Sensitivity")
    sens["A1"] = "Equity value per share (USD): WACC (rows) by terminal growth (columns)"
    sens["A1"].font = Font(bold=True, size=13)
    shifts = [-0.01, -0.005, 0.0, 0.005, 0.01]
    growths = [0.01, 0.015, 0.02, 0.025, 0.03]
    sens["A3"] = "WACC \\ g"
    for j, g in enumerate(growths):
        sens.cell(row=3, column=2 + j, value=g).number_format = "0.0%"
        sens.cell(row=3, column=2 + j).font = INPUT_FONT
    for i, shift in enumerate(shifts):
        r = 4 + i
        sens.cell(row=r, column=1, value=f"=DCF!$B$17+({shift})").number_format = "0.00%"
        for j in range(5):
            gcell = f"{get_column_letter(2 + j)}$3"
            w = f"$A{r}"
            pv = f"SUMPRODUCT(DCF!$C$32:$G$32,1/(1+{w})^(DCF!$C$22:$G$22-0.5*DCF!$B$19))"
            tv = f"(DCF!$G$27*(1+{gcell})-DCF!$B$8*DCF!$G$25*{gcell})/({w}-{gcell})/(1+{w})^(DCF!$G$22-0.5*DCF!$B$19)"
            sens.cell(row=r, column=2 + j, value=f"=({pv}+{tv}+DCF!$B$54)/DCF!$B$18").number_format = "0.0"
    sens.column_dimensions["A"].width = 14

    # ---------------------------------------------------------- Checks
    chk = wb.create_sheet("Checks")
    chk["A1"] = "Checks (every value should read 0 or TRUE)"
    chk["A1"].font = Font(bold=True, size=13)
    r = 3
    for company in COMPANIES:
        for year in YEARS:
            chk[f"A{r}"] = f"Balance sheet identity {company} {year}: assets - equity - liabilities"
            chk[f"B{r}"] = (
                f"={sumifs(n, company, year, 'total_assets')}-{sumifs(n, company, year, 'total_equity')}-{sumifs(n, company, year, 'total_liabilities')}"
            )
            r += 1
    chk[f"A{r}"] = "WACC exceeds terminal growth"
    chk[f"B{r}"] = "=DCF!B17>DCF!B9"
    r += 1
    chk[f"A{r}"] = "Bridge sums to equity value"
    chk[f"B{r}"] = "=ABS(SUM(DCF!B45:B52)-DCF!B53)<0.001"
    chk.column_dimensions["A"].width = 80

    OUT.parent.mkdir(parents=True, exist_ok=True)
    for sheet in wb:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.alignment is None or not cell.alignment.wrap_text:
                    cell.alignment = Alignment(vertical="top")
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
